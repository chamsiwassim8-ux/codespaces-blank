from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import streamlit as st


WORKBOOK_PATH = Path(__file__).with_name("Container infos.xlsx")
PRODUCT_SHEET = "Product Specs"
LOAD_SHEET = "Container Load"


@dataclass(frozen=True)
class Product:
    description: str
    sap_code: str
    customer_code: str
    mass_g: int | None
    size: str | None

    @property
    def default_spoon_cap(self) -> bool:
        return bool(re.search(r"ultima|ultimate", self.description, re.I))

    @property
    def label(self) -> str:
        return (
            f"{self.description} | SAP: {self.sap_code or '-'}"
            f" | Client: {self.customer_code or '-'}"
        )


@dataclass(frozen=True)
class Capacity:
    tins_per_unit: int
    tins_per_pallet: int
    pallets_per_unit: int | None
    unit_name: str


@dataclass(frozen=True)
class ShipmentLine:
    product: Product
    quantity: int
    with_spoon_cap: bool
    capacity: Capacity


def _number(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(float(str(value).replace(",", ".")))
    except ValueError:
        return None


def detect_mass(description: str) -> int | None:
    match = re.search(r"(?<!\d)(400|800|900)\s*(?:g|gr|gram(?:me)?s?)?\b", description, re.I)
    return int(match.group(1)) if match else None


def detect_size(description: str) -> str | None:
    mass = detect_mass(description)
    if mass is None:
        return None
    return "400" if mass == 400 else "800/900"


def load_workbook_data(file_obj: Any) -> tuple[list[Product], Any]:
    workbook = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    if PRODUCT_SHEET not in workbook.sheetnames or LOAD_SHEET not in workbook.sheetnames:
        raise ValueError("The workbook must contain 'Product Specs' and 'Container Load' sheets.")

    product_sheet = workbook[PRODUCT_SHEET]
    products: list[Product] = []
    for row in product_sheet.iter_rows(min_row=2, values_only=True):
        description = str(row[8] or "").strip() if len(row) > 8 else ""
        if not description:
            continue
        products.append(
            Product(
                description=description,
                sap_code=str(row[6] or "").strip() if len(row) > 6 else "",
                customer_code=str(row[4] or "").strip() if len(row) > 4 else "",
                mass_g=detect_mass(description),
                size=detect_size(description),
            )
        )

    if not products:
        raise ValueError("No products with a description were found in 'Product Specs'.")
    return products, workbook[LOAD_SHEET]


def _cell_number(sheet: Any, row: int, column: int, fallback: tuple[int, int] | None = None) -> int:
    value = _number(sheet.cell(row=row, column=column).value)
    if value is None and fallback is not None:
        value = _number(sheet.cell(row=fallback[0], column=fallback[1]).value)
    if value is None:
        raise ValueError(f"No capacity is defined at {sheet.cell(row=row, column=column).coordinate}.")
    return value


def get_capacity(
    sheet: Any,
    mode: str,
    size: str,
    with_spoon_cap: bool,
    packing: str | None = None,
    container_size: str | None = None,
    container_type: str | None = None,
    truck_type: str | None = None,
) -> Capacity:
    size_index = 0 if size == "400" else 1
    if mode == "Container":
        columns = {
            ("20'", "Reefer"): (3, 4),
            ("20'", "DV"): (5, 6),
            ("40'", "Reefer"): (7, 8),
            ("40'", "DV"): (9, 10),
            ("40'", "HC"): (11, 12),
        }
        pair = columns.get((container_size, container_type))
        if pair is None:
            raise ValueError("This container size and type combination is not available.")
        column = pair[size_index]
        if packing == "Non Palletized":
            base_row = 7 if with_spoon_cap else 8
            tins = _cell_number(sheet, base_row, column, (7, column) if not with_spoon_cap else None)
            return Capacity(tins, 0, None, "containers")
        if packing != "Palletized":
            raise ValueError("Choose a valid container packing type.")
        base_row = 15 if with_spoon_cap else 16
        tins_per_pallet = _cell_number(sheet, base_row, column, (15, column) if not with_spoon_cap else None)
        pallets = _number(sheet.cell(row=14, column=column).value)
        if pallets is None:
            pallets = _number(sheet.cell(row=14, column=pair[0]).value)
        if pallets is None:
            raise ValueError("No pallet capacity is defined for this container.")
        total_row = 17 if with_spoon_cap else 18
        tins = _number(sheet.cell(row=total_row, column=column).value)
        if tins is None:
            tins = pallets * tins_per_pallet
        return Capacity(tins, tins_per_pallet, pallets, "containers")

    if mode == "Truck":
        if truck_type == "Dry Truck":
            pair = (3, 4)
        elif truck_type == "Dry Truck (Iraq)":
            pair = (5, 6)
        else:
            raise ValueError("Choose a valid truck type.")
        column = pair[size_index]
        tins_per_pallet_row = 24 if with_spoon_cap else 25
        total_row = 26 if with_spoon_cap else 27
        tins_per_pallet = _cell_number(
            sheet,
            tins_per_pallet_row,
            column,
            (24, column) if not with_spoon_cap else None,
        )
        tins = _cell_number(sheet, total_row, column, (26, column) if not with_spoon_cap else None)
        pallets = _cell_number(sheet, 23, column)
        return Capacity(tins, tins_per_pallet, pallets, "trucks")

    if mode == "Airshipment":
        column = 10 if size_index == 0 else 11
        tins_row = 26 if with_spoon_cap else 27
        tins = _cell_number(sheet, tins_row, column, (26, column) if not with_spoon_cap else None)
        return Capacity(tins, tins, 1, "airloads")

    raise ValueError("Choose a valid shipment mode.")


def calculate_shipment(
    lines: list[ShipmentLine],
    mode: str,
    packing: str | None,
) -> dict[str, Any]:
    if not lines:
        raise ValueError("Add at least one product to the shipment.")

    palletized = mode == "Truck" or (mode == "Container" and packing == "Palletized")
    if mode == "Airshipment":
        grouped_airloads: dict[tuple[int | None, str | None, bool], dict[str, int | Capacity]] = {}
        for line in lines:
            key = (line.product.mass_g, line.product.size, line.with_spoon_cap)
            if key not in grouped_airloads:
                grouped_airloads[key] = {"quantity": 0, "capacity": line.capacity}
            grouped_airloads[key]["quantity"] = int(grouped_airloads[key]["quantity"]) + line.quantity

        airload_groups = [
            (
                math.ceil(int(group["quantity"]) / group["capacity"].tins_per_unit),
                int(group["quantity"]),
                group["capacity"],
            )
            for group in grouped_airloads.values()
        ]
        units = sum(group[0] for group in airload_groups)
        last_group_units, last_group_quantity, last_group_capacity = airload_groups[-1]
        tins_in_last_airload = last_group_quantity - last_group_capacity.tins_per_unit * (last_group_units - 1)
        utilization = tins_in_last_airload / last_group_capacity.tins_per_unit * 100
        remaining_tins = last_group_capacity.tins_per_unit - tins_in_last_airload
        pallets = units
        pallet_positions = None
        remaining_pallet_positions = None
        load_share = None
        mixed_mass = False
    elif palletized:
        grouped: dict[tuple[int | None, str | None, bool], dict[str, int | Capacity]] = {}
        for line in lines:
            if line.capacity.tins_per_pallet <= 0 or line.capacity.pallets_per_unit is None:
                raise ValueError("Pallet capacity is missing for one of the selected products.")
            key = (line.product.mass_g, line.product.size, line.with_spoon_cap)
            if key not in grouped:
                grouped[key] = {"quantity": 0, "capacity": line.capacity}
            grouped[key]["quantity"] = int(grouped[key]["quantity"]) + line.quantity

        pallets = sum(
            math.ceil(int(group["quantity"]) / group["capacity"].tins_per_pallet)
            for group in grouped.values()
        )
        mixed_mass = mode == "Container" and packing == "Palletized" and len(
            {line.product.mass_g for line in lines}
        ) > 1
        pallet_positions = (
            max(group["capacity"].pallets_per_unit for group in grouped.values())
            if mixed_mass
            else min(group["capacity"].pallets_per_unit for group in grouped.values())
        )
        if mixed_mass:
            pallet_positions -= 1
        if pallet_positions <= 0:
            raise ValueError("The available pallet capacity is too small for this product mix.")

        units = math.ceil(pallets / pallet_positions)
        pallets_in_last_unit = pallets - pallet_positions * (units - 1)
        utilization = pallets_in_last_unit / pallet_positions * 100
        remaining_pallet_positions = pallet_positions - pallets_in_last_unit
        load_share = None
        remaining_tins = None
    else:
        load_share = sum(line.quantity / line.capacity.tins_per_unit for line in lines)
        units = math.ceil(load_share)
        utilization = (load_share - (units - 1)) * 100
        remaining_tins = (
            lines[0].capacity.tins_per_unit
            - (lines[0].quantity - lines[0].capacity.tins_per_unit * (units - 1))
            if len(lines) == 1
            else None
        )
        remaining_pallet_positions = None
        pallet_positions = None
        pallets = 0
        mixed_mass = False

    return {
        "units": units,
        "pallets": pallets,
        "pallet_positions": pallet_positions,
        "remaining_pallet_positions": remaining_pallet_positions,
        "remaining_tins": remaining_tins,
        "load_share": load_share,
        "last_unit_utilization": utilization,
        "mixed_mass": mixed_mass,
        "palletized": palletized,
    }


def distribute_containers(
    lines: list[ShipmentLine],
    packing: str,
    pallet_positions: int | None,
) -> list[dict[str, Any]]:
    containers: list[dict[str, Any]] = []

    def new_container() -> dict[str, Any]:
        return {"items": {}, "total_tins": 0, "capacity_share": 0.0, "pallets": []}

    def add_tins(container: dict[str, Any], line_index: int, quantity: int) -> None:
        line = lines[line_index]
        item = container["items"].setdefault(
            line_index,
            {
                "Product": line.product.description,
                "Mass (g)": line.product.mass_g,
                "Spoon cap": "with SpC" if line.with_spoon_cap else "without SpC",
                "Tins": 0,
            },
        )
        item["Tins"] += quantity
        container["total_tins"] += quantity

    if packing == "Non Palletized":
        remaining = [line.quantity for line in lines]
        while any(remaining):
            container = new_container()
            for line_index, line in enumerate(lines):
                if remaining[line_index] <= 0:
                    continue
                capacity = line.capacity.tins_per_unit
                available_share = max(0.0, 1.0 - container["capacity_share"])
                tins_that_fit = math.floor(available_share * capacity + 1e-9)
                if tins_that_fit <= 0:
                    continue
                quantity = min(remaining[line_index], tins_that_fit)
                add_tins(container, line_index, quantity)
                container["capacity_share"] += quantity / capacity
                remaining[line_index] -= quantity

            if not container["items"]:
                raise ValueError("Unable to fit any remaining tins into a container.")
            container["max_capacity_tins"] = container["total_tins"] / container["capacity_share"]
            container["fill_rate"] = min(100.0, container["capacity_share"] * 100)
            container["pallets_used"] = None
            container["pallets_max"] = None
            containers.append(container)
        return containers

    if pallet_positions is None or pallet_positions <= 0:
        raise ValueError("Pallet capacity is missing for this container.")

    pallet_groups: dict[tuple[int | None, str | None, bool], list[int]] = {}
    for line_index, line in enumerate(lines):
        pallet_groups.setdefault(
            (line.product.mass_g, line.product.size, line.with_spoon_cap), []
        ).append(line_index)

    pallets_to_pack = []
    for group_lines in pallet_groups.values():
        pallet_capacity = lines[group_lines[0]].capacity.tins_per_pallet
        remaining = {line_index: lines[line_index].quantity for line_index in group_lines}
        while any(remaining.values()):
            pallet = {"capacity_tins": pallet_capacity, "total_tins": 0, "items": []}
            open_tins = pallet_capacity
            for line_index in group_lines:
                if open_tins <= 0:
                    break
                quantity = min(remaining[line_index], open_tins)
                if quantity <= 0:
                    continue
                pallet["items"].append((line_index, quantity))
                pallet["total_tins"] += quantity
                remaining[line_index] -= quantity
                open_tins -= quantity
            pallets_to_pack.append(pallet)

    max_tins_per_pallet = max(line.capacity.tins_per_pallet for line in lines)
    for pallet in pallets_to_pack:
        if not containers or len(containers[-1]["pallets"]) >= pallet_positions:
            containers.append(new_container())
        container = containers[-1]
        container["pallets"].append(pallet)
        for line_index, quantity in pallet["items"]:
            add_tins(container, line_index, quantity)

    for container in containers:
        pallets_used = len(container["pallets"])
        max_capacity_tins = sum(pallet["capacity_tins"] for pallet in container["pallets"])
        max_capacity_tins += (pallet_positions - pallets_used) * max_tins_per_pallet
        container["max_capacity_tins"] = max_capacity_tins
        container["fill_rate"] = container["total_tins"] / max_capacity_tins * 100
        container["pallets_used"] = pallets_used
        container["pallets_max"] = pallet_positions

    return containers


def _load_source(uploaded_file: Any) -> tuple[list[Product], Any]:
    source = uploaded_file if uploaded_file is not None else WORKBOOK_PATH
    return load_workbook_data(source)


def _reset_spoon_cap(line_id: int) -> None:
    product = st.session_state[f"product_{line_id}"]
    st.session_state[f"spoon_cap_{line_id}"] = (
        "with SpC" if product.default_spoon_cap else "without SpC"
    )


def _resolve_product(identifier: str, products: list[Product]) -> Product:
    query = " ".join(identifier.split()).casefold()
    if not query:
        raise ValueError("Enter a product description, SAP code, or customer code.")

    def unique_matches(matches: list[Product], include_sap: bool = False) -> list[Product]:
        return list(
            {
                (
                    " ".join(product.description.split()).casefold(),
                    product.sap_code.casefold() if include_sap else "",
                    product.mass_g,
                    product.size,
                ): product
                for product in matches
            }.values()
        )

    exact_description = unique_matches(
        [product for product in products if query == " ".join(product.description.split()).casefold()]
    )
    if len(exact_description) == 1:
        return exact_description[0]
    if len(exact_description) > 1:
        raise ValueError(f"'{identifier}' matches multiple product descriptions.")

    exact_sap = unique_matches(
        [product for product in products if query == product.sap_code.casefold()],
        include_sap=True,
    )
    if len(exact_sap) == 1:
        return exact_sap[0]
    if len(exact_sap) > 1:
        raise ValueError(f"'{identifier}' matches multiple products; use a full description.")

    exact_customer_code = unique_matches(
        [product for product in products if query == product.customer_code.casefold()],
        include_sap=True,
    )
    if len(exact_customer_code) == 1:
        return exact_customer_code[0]
    if len(exact_customer_code) > 1:
        raise ValueError(f"'{identifier}' matches multiple products; use a SAP code or full description.")

    exact_label = unique_matches(
        [product for product in products if query == product.label.casefold()],
        include_sap=True,
    )
    if len(exact_label) == 1:
        return exact_label[0]
    if len(exact_label) > 1:
        raise ValueError(f"'{identifier}' matches multiple products; use a SAP code or full description.")

    description_matches = unique_matches([
        product for product in products if query in " ".join(product.description.split()).casefold()
    ])
    if len(description_matches) == 1:
        return description_matches[0]
    if len(description_matches) > 1:
        raise ValueError(f"'{identifier}' matches multiple products; use a SAP code or full description.")
    raise ValueError(f"No product matches '{identifier}'.")


st.set_page_config(page_title="Shipment planner", page_icon="C", layout="wide")
st.markdown(
    """
    <style>
    :root { --ink: #172b2b; --muted: #586b67; --line: #d6e1db; --green: #176b52; }
    .stApp { background: linear-gradient(125deg, #f5f8f2 0%, #eef5ef 48%, #f8f6ec 100%); color: var(--ink); }
    .stApp, .stApp p, .stApp label, .stApp [data-testid="stCaptionContainer"] { color: var(--ink); }
    .stApp input, .stApp textarea, .stApp [data-baseweb="select"] * { color: var(--ink) !important; }
    .stApp input, .stApp textarea { background-color: #ffffff !important; }
    .stApp [data-baseweb="popover"] *, .stApp [role="option"] { color: var(--ink) !important; }
    .stApp [data-testid="stDataEditor"] { color: var(--ink); }
    [data-testid="stHeader"] { background: transparent; }
    .block-container { max-width: 1180px; padding-top: 2rem; padding-bottom: 3rem; }
    h1, h2, h3 { color: var(--ink); font-family: Georgia, "Times New Roman", serif; letter-spacing: 0; }
    h1 { font-size: 2.5rem; margin-bottom: .25rem; }
    .eyebrow { color: var(--green); font-size: .75rem; font-weight: 700; letter-spacing: 0; text-transform: uppercase; }
    .subhead { color: var(--muted); margin: 0 0 1.4rem; }
    div[data-testid="stMetric"] { background: rgba(255,255,255,.72); border: 1px solid var(--line); border-radius: 6px; padding: 1rem 1.1rem; }
    div[data-testid="stMetricLabel"] { color: var(--muted); }
    div[data-testid="stMetricValue"] { color: var(--green); }
    div[data-testid="stTabs"] button { color: var(--muted); }
    div[data-testid="stTabs"] button[aria-selected="true"] { color: var(--green); }
    hr { border-color: var(--line); }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown('<div class="eyebrow">Logistics workspace / Capacity planning</div>', unsafe_allow_html=True)
st.title("Shipment planner")
st.markdown('<p class="subhead">Build a load plan from your product and shipment requirements.</p>', unsafe_allow_html=True)

with st.sidebar:
    st.subheader("Data source")
    uploaded = st.file_uploader("Excel workbook", type=["xlsx"], help="Defaults to Container infos.xlsx in the app folder.")

try:
    products, load_sheet = _load_source(uploaded)
except (OSError, ValueError, KeyError, openpyxl.utils.exceptions.InvalidFileException) as exc:
    st.error(f"Unable to read the workbook: {exc}")
    st.stop()

unrecognized = sum(product.size is None for product in products)
if unrecognized:
    st.sidebar.caption(f"{unrecognized} product description(s) do not include a supported size (400g, 800g or 900g).")

left, right = st.columns([1.1, .9], gap="large")
product_specs = []
with left:
    st.subheader("1. Products and quantities")
    entry_method = st.radio(
        "Product entry",
        ["Individual lines", "Paste table"],
        horizontal=True,
        key="product_entry_method",
    )
    if entry_method == "Paste table":
        if "bulk_product_table" not in st.session_state:
            st.session_state.bulk_product_table = pd.DataFrame(
                {
                    "Product identifier": [""] * 8,
                    "Quantity (tins)": pd.Series([None] * 8, dtype="Int64"),
                    "SpC override": ["Auto"] * 8,
                }
            )
        edited_table = st.data_editor(
            st.session_state.bulk_product_table,
            key="bulk_product_editor",
            num_rows="dynamic",
            width="stretch",
            hide_index=True,
            column_config={
                "Product identifier": st.column_config.TextColumn(),
                "Quantity (tins)": st.column_config.NumberColumn(min_value=1, step=1, format="%d"),
                "SpC override": st.column_config.SelectboxColumn(
                    options=["Auto", "with SpC", "without SpC"],
                    default="Auto",
                ),
            },
        )
        st.session_state.bulk_product_table = edited_table
        table_errors = []
        for row_number, row in enumerate(edited_table.to_dict("records"), start=1):
            raw_identifier = row.get("Product identifier")
            identifier = "" if pd.isna(raw_identifier) else str(raw_identifier).strip()
            raw_quantity = row.get("Quantity (tins)")
            quantity_is_blank = pd.isna(raw_quantity) or str(raw_quantity).strip() == ""
            if not identifier and quantity_is_blank:
                continue
            if not identifier:
                table_errors.append(f"Row {row_number}: product identifier is missing.")
                continue
            if quantity_is_blank:
                table_errors.append(f"Row {row_number}: quantity is missing.")
                continue
            try:
                quantity_value = float(raw_quantity)
                if quantity_value < 1 or not quantity_value.is_integer():
                    raise ValueError
                selected_product = _resolve_product(identifier, products)
            except ValueError as exc:
                table_errors.append(f"Row {row_number}: {exc or 'quantity must be a positive whole number.'}")
                continue

            spoon_cap_value = str(row.get("SpC override") or "Auto").strip().casefold()
            if spoon_cap_value in {"", "auto"}:
                with_spoon_cap = selected_product.default_spoon_cap
            elif spoon_cap_value == "with spc":
                with_spoon_cap = True
            elif spoon_cap_value == "without spc":
                with_spoon_cap = False
            else:
                table_errors.append(f"Row {row_number}: SpC must be Auto, with SpC, or without SpC.")
                continue
            product_specs.append((selected_product, int(quantity_value), with_spoon_cap))

        if table_errors:
            st.error("  ".join(table_errors))
            st.stop()
        if not product_specs:
            st.warning("Enter at least one product and quantity.")
            st.stop()
    else:
        if "shipment_line_ids" not in st.session_state:
            st.session_state.shipment_line_ids = [0]
            st.session_state.next_shipment_line_id = 1
        for position, line_id in enumerate(list(st.session_state.shipment_line_ids), start=1):
            product_key = f"product_{line_id}"
            quantity_key = f"quantity_{line_id}"
            spoon_cap_key = f"spoon_cap_{line_id}"
            st.session_state.setdefault(product_key, products[0])
            st.session_state.setdefault(quantity_key, 10000)
            st.session_state.setdefault(
                spoon_cap_key,
                "with SpC" if st.session_state[product_key].default_spoon_cap else "without SpC",
            )

            product_column, quantity_column, spoon_cap_column, remove_column = st.columns(
                [4.3, 1.4, 1.8, .8], vertical_alignment="bottom"
            )
            with product_column:
                selected_product = st.selectbox(
                    f"Product {position}",
                    products,
                    format_func=lambda product: product.label,
                    key=product_key,
                    on_change=_reset_spoon_cap,
                    args=(line_id,),
                )
            with quantity_column:
                quantity = st.number_input("Tins", min_value=1, step=1000, key=quantity_key)
            with spoon_cap_column:
                spoon_cap = st.selectbox("Spoon cap", ["with SpC", "without SpC"], key=spoon_cap_key)
            if len(st.session_state.shipment_line_ids) > 1 and remove_column.button(
                "Remove", key=f"remove_{line_id}"
            ):
                st.session_state.shipment_line_ids.remove(line_id)
                st.rerun()
            st.caption(
                f"Mass: {selected_product.mass_g or 'Unknown'} g  |  "
                f"SAP: {selected_product.sap_code or '-'}  |  "
                f"Customer code: {selected_product.customer_code or '-'}  |  "
                f"Auto SpC: {'with SpC' if selected_product.default_spoon_cap else 'without SpC'}"
            )
            product_specs.append((selected_product, int(quantity), spoon_cap == "with SpC"))

        if st.button("Add another product", type="secondary"):
            line_id = st.session_state.next_shipment_line_id
            st.session_state.next_shipment_line_id += 1
            st.session_state.shipment_line_ids.append(line_id)
            st.session_state[f"product_{line_id}"] = products[0]
            st.session_state[f"quantity_{line_id}"] = 10000
            st.session_state[f"spoon_cap_{line_id}"] = (
                "with SpC" if products[0].default_spoon_cap else "without SpC"
            )
            st.rerun()

with right:
    st.subheader("2. Shipment configuration")
    mode = st.selectbox("Shipment mode", ["Container", "Truck", "Airshipment"])
    packing = None
    container_size = None
    container_type = None
    truck_type = None
    if mode == "Container":
        packing = st.selectbox("Packing type", ["Palletized", "Non Palletized"])
        container_size = st.selectbox("Container size", ["20'", "40'"])
        allowed_types = ["Reefer", "DV"] if container_size == "20'" else ["Reefer", "DV", "HC"]
        container_type = st.selectbox("Container type", allowed_types)
    elif mode == "Truck":
        truck_type = st.selectbox("Truck type", ["Dry Truck", "Dry Truck (Iraq)"])
    else:
        st.text_input("Airshipment type", value="Standard Airload", disabled=True)

shipment_lines = []
for product, quantity, with_spoon_cap in product_specs:
    if product.size is None:
        st.warning(
            f"Size could not be detected for {product.description}. "
            "Select a product whose description includes 400g, 800g or 900g."
        )
        st.stop()
    try:
        capacity = get_capacity(
            load_sheet,
            mode=mode,
            size=product.size,
            with_spoon_cap=with_spoon_cap,
            packing=packing,
            container_size=container_size,
            container_type=container_type,
            truck_type=truck_type,
        )
    except ValueError as exc:
        st.error(str(exc))
        st.stop()
    shipment_lines.append(
        ShipmentLine(
            product=product,
            quantity=quantity,
            with_spoon_cap=with_spoon_cap,
            capacity=capacity,
        )
    )

try:
    result = calculate_shipment(shipment_lines, mode, packing)
except ValueError as exc:
    st.error(str(exc))
    st.stop()

st.divider()
st.subheader("Shipment summary")
st.caption(f"{len(shipment_lines)} product line(s)  |  {mode}")
if mode == "Container":
    st.caption(f"{packing}  |  {container_size} {container_type}")
elif mode == "Truck":
    st.caption(truck_type)
else:
    st.caption("Standard Airload")

unit_label = {"containers": "Containers required", "trucks": "Trucks required", "airloads": "Airloads required"}[capacity.unit_name]
metric_columns = st.columns(4)
capacity_categories = {
    (line.product.mass_g, line.product.size, line.with_spoon_cap) for line in shipment_lines
}
if mode == "Container" and packing == "Palletized":
    capacity_label = "Pallet positions per load"
    capacity_value = f"{result['pallet_positions']:,}"
else:
    capacity_label = {
        "Container": "Tins per container",
        "Truck": "Tins per truck",
        "Airshipment": "Tins per airload",
    }[mode]
    capacity_value = (
        f"{shipment_lines[0].capacity.tins_per_unit:,}"
        if len(capacity_categories) == 1
        else "Varies by product"
    )
metric_columns[0].metric(capacity_label, capacity_value)
metric_columns[1].metric(unit_label, f"{result['units']:,}")
metric_columns[2].metric("Pallets required", f"{result['pallets']:,}" if result["pallets"] else "Not palletized")
metric_columns[3].metric("Last unit utilization", f"{result['last_unit_utilization']:.1f}%")

# --- PROMINENT 20' CONTAINER RECOMMENDATION BUBBLE ---
if mode == "Container" and container_size == "40'" and result["last_unit_utilization"] < 80.0:
    try:
        lines_20ft = []
        for line in shipment_lines:
            cap_20 = get_capacity(
                load_sheet,
                mode="Container",
                size=line.product.size,
                with_spoon_cap=line.with_spoon_cap,
                packing=packing,
                container_size="20'",
                container_type=container_type,
            )
            lines_20ft.append(
                ShipmentLine(
                    product=line.product,
                    quantity=line.quantity,
                    with_spoon_cap=line.with_spoon_cap,
                    capacity=cap_20,
                )
            )
        res_20 = calculate_shipment(lines_20ft, mode="Container", packing=packing)
        
        st.warning(
            f"📦 **Container Optimization Suggestion:** Your last 40' container is under-utilized "
            f"at **{result['last_unit_utilization']:.1f}%**.\n\n"
            f"💡 Switching to **20' {container_type}** containers would require **{res_20['units']} container(s)** "
            f"with a revised last-unit utilization of **{res_20['last_unit_utilization']:.1f}%**.",
            icon="💡"
        )
    except Exception as e:
        # Temporarily show if there's a lookup issue in the Excel grid
        st.caption(f"*(Note: 20' alternative comparison unavailable for this specific configuration: {e})*")

line_summary = []
for line in shipment_lines:
    if result["palletized"]:
        tins_per_load_unit = line.capacity.tins_per_pallet * result["pallet_positions"]
    else:
        tins_per_load_unit = line.capacity.tins_per_unit
    line_summary.append(
        {
            "Product": line.product.description,
            "Mass (g)": line.product.mass_g,
            "Quantity (tins)": line.quantity,
            "Auto SpC": "with SpC" if line.product.default_spoon_cap else "without SpC",
            "Selected SpC": "with SpC" if line.with_spoon_cap else "without SpC",
            "Tins / pallet": line.capacity.tins_per_pallet or "-",
            "Tins / load unit": tins_per_load_unit,
        }
    )
st.dataframe(line_summary, hide_index=True, width="stretch")

if result["mixed_mass"]:
    st.info(
        "Mixed masses use separate pallets. One pallet position is deducted from each "
        "palletized container's capacity."
    )

if result["palletized"]:
    remaining_label = "Pallet positions open in the last load"
    remaining_value = f"{result['remaining_pallet_positions']:,}"
elif result["remaining_tins"] is not None:
    remaining_label = "Unfilled capacity in the last load"
    remaining_value = f"{result['remaining_tins']:,} tins"
else:
    remaining_label = "Unused combined capacity in the last load"
    remaining_value = f"{result['units'] - result['load_share']:.2f} load-equivalent"
st.markdown(f"**{remaining_label}:** {remaining_value}")
st.progress(result["last_unit_utilization"] / 100, text=f"Last load: {result['last_unit_utilization']:.1f}% utilized")

if mode == "Container":
    try:
        container_plan = distribute_containers(
            shipment_lines,
            packing,
            result["pallet_positions"],
        )
    except ValueError as exc:
        st.error(str(exc))
        st.stop()

    st.divider()
    st.subheader("Container-by-container plan")
    for container_number, container in enumerate(container_plan, start=1):
        fill_rate = container["fill_rate"]
        status = "Full" if fill_rate >= 99.95 else "Partially filled"
        with st.container(border=True):
            st.markdown(f"**Container #{container_number} · {container_size} {container_type} · {status}**")
            detail_columns = st.columns(3)
            detail_columns[0].metric("Fill rate", f"{fill_rate:.1f}%")
            detail_columns[1].metric("Total tins", f"{container['total_tins']:,}")
            if container["pallets_used"] is not None:
                detail_columns[2].metric("Pallets used", f"{container['pallets_used']} / {container['pallets_max']}")
            else:
                detail_columns[2].metric("Max capacity", f"{int(container['max_capacity_tins']):,} tins")
            st.progress(fill_rate / 100, text=f"Utilization: {fill_rate:.1f}%")

            item_list = list(container["items"].values())
            st.dataframe(item_list, hide_index=True, width="stretch")
